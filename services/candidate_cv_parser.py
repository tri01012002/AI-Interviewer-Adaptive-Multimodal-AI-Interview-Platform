"""Bounded, deterministic CV extraction for PDF and DOCX documents."""

from __future__ import annotations

import re
import zipfile
from io import BytesIO
from pathlib import PurePosixPath

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from lxml.etree import XMLSyntaxError
from pydantic import BaseModel, ConfigDict, Field
from pypdf import PdfReader
from pypdf.errors import PdfReadError

PARSER_VERSION = "candidate-cv-heuristic-v1"
MAX_DOCUMENT_PAGES = 100
MAX_EXTRACTED_CHARS = 200_000
MAX_DOCX_ENTRIES = 500
MAX_DOCX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024
UNTRUSTED_INSTRUCTION_PATTERN = re.compile(
    r"\b(?:ignore|disregard|override)\b.{0,80}\b(?:instructions?|system prompt|criteria)\b"
    r"|\b(?:system prompt|give the candidate|award the candidate|assign .* score|"
    r"rate the candidate|do not follow)\b",
    flags=re.IGNORECASE,
)
ALLOWED_MEDIA_TYPES = {
    ".pdf": {"application/pdf", "application/octet-stream"},
    ".docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/octet-stream",
    },
}


class CVDocumentError(ValueError):
    code = "invalid_cv_document"


class EmptyCVDocumentError(CVDocumentError):
    code = "empty_cv_document"


class ParsedEvidence(BaseModel):
    model_config = ConfigDict(extra="forbid")

    text: str
    start_offset: int
    end_offset: int


class ParsedCandidateCV(BaseModel):
    model_config = ConfigDict(extra="forbid")

    parser_version: str
    employment_history: list[ParsedEvidence] = Field(default_factory=list)
    skills_and_technologies: list[ParsedEvidence] = Field(default_factory=list)
    education_and_qualifications: list[ParsedEvidence] = Field(default_factory=list)
    project_experience: list[ParsedEvidence] = Field(default_factory=list)


def validate_cv_upload(filename: str, media_type: str | None, content: bytes) -> str:
    extension = PurePosixPath(filename.replace("\\", "/")).suffix.lower()
    if extension not in ALLOWED_MEDIA_TYPES:
        raise CVDocumentError("Only PDF and DOCX CVs are accepted")
    normalized_media_type = (media_type or "").split(";", 1)[0].strip().lower()
    if normalized_media_type not in ALLOWED_MEDIA_TYPES[extension]:
        raise CVDocumentError("The supplied media type does not match an accepted CV format")
    if extension == ".pdf":
        _extract_pdf_text(content, validate_only=True)
    else:
        _extract_docx_text(content, validate_only=True)
    return extension


def parse_candidate_cv(content: bytes, extension: str) -> tuple[str, ParsedCandidateCV]:
    if extension == ".pdf":
        extracted_text = _extract_pdf_text(content)
    elif extension == ".docx":
        extracted_text = _extract_docx_text(content)
    else:
        raise CVDocumentError("Unsupported CV format")

    extracted_text = extracted_text[:MAX_EXTRACTED_CHARS]
    if not extracted_text.strip():
        raise EmptyCVDocumentError("The CV contains no extractable text")
    return extracted_text, _extract_structured_fields(extracted_text)


def _extract_pdf_text(content: bytes, *, validate_only: bool = False) -> str:
    if not content.startswith(b"%PDF-"):
        raise CVDocumentError("File content is not a PDF document")
    try:
        reader = PdfReader(BytesIO(content), strict=True)
        if reader.is_encrypted:
            raise CVDocumentError("Encrypted PDF CVs are not supported")
        if len(reader.pages) > MAX_DOCUMENT_PAGES:
            raise CVDocumentError("PDF exceeds the maximum supported page count")
        if validate_only:
            return ""
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except (PdfReadError, ValueError) as exc:
        raise CVDocumentError("PDF CV is malformed or cannot be parsed") from exc


def _extract_docx_text(content: bytes, *, validate_only: bool = False) -> str:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_ENTRIES:
                raise CVDocumentError("DOCX contains too many archive entries")
            if sum(entry.file_size for entry in entries) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise CVDocumentError("DOCX expands beyond the supported document size")
            for entry in entries:
                archive_path = PurePosixPath(entry.filename)
                if archive_path.is_absolute() or ".." in archive_path.parts:
                    raise CVDocumentError("DOCX contains an unsafe archive path")
            if "[Content_Types].xml" not in archive.namelist() or "word/document.xml" not in archive.namelist():
                raise CVDocumentError("DOCX package is missing required document parts")
            if archive.testzip() is not None:
                raise CVDocumentError("DOCX package contains a corrupt archive entry")
        if validate_only:
            return ""
        document = Document(BytesIO(content))
        blocks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            blocks.extend(cell.text for row in table.rows for cell in row.cells)
        return "\n".join(blocks)
    except CVDocumentError:
        raise
    except (zipfile.BadZipFile, PackageNotFoundError, XMLSyntaxError, KeyError, ValueError) as exc:
        raise CVDocumentError("DOCX CV is malformed or cannot be parsed") from exc


def _excerpt(text: str, start: int, end: int) -> ParsedEvidence:
    left = max(0, text.rfind("\n", 0, start) + 1)
    right = text.find("\n", end)
    if right < 0:
        right = min(len(text), end + 180)
    else:
        right = min(right, left + 300)
    return ParsedEvidence(
        text=text[left:right].strip()[:300],
        start_offset=left,
        end_offset=right,
    )


def _extract_structured_fields(text: str) -> ParsedCandidateCV:
    lines = [
        (match.group().strip(), match.start(), match.end())
        for match in re.finditer(r"[^\r\n]+", text)
        if match.group().strip()
    ]
    evidence: dict[str, list[ParsedEvidence]] = {
        "employment_history": [],
        "skills_and_technologies": [],
        "education_and_qualifications": [],
        "project_experience": [],
    }

    skill_terms = (
        "python", "java", "javascript", "typescript", "c++", "c#", "go", "rust",
        "sql", "postgresql", "mysql", "mongodb", "redis", "fastapi", "django",
        "flask", "react", "node.js", "aws", "azure", "gcp", "docker", "kubernetes",
        "terraform", "git", "linux", "pytorch", "tensorflow", "scikit-learn",
    )
    for term in skill_terms:
        pattern = re.compile(
            rf"(?<![\w+#.]){re.escape(term)}(?![\w+#.])",
            flags=re.IGNORECASE,
        )
        for match in pattern.finditer(text):
            line_start = max(0, text.rfind("\n", 0, match.start()) + 1)
            line_end = text.find("\n", match.end())
            if line_end < 0:
                line_end = len(text)
            if UNTRUSTED_INSTRUCTION_PATTERN.search(text[line_start:line_end]):
                continue
            evidence["skills_and_technologies"].append(
                ParsedEvidence(
                    text=match.group(),
                    start_offset=match.start(),
                    end_offset=match.end(),
                )
            )

    date_pattern = re.compile(
        r"\b(?:19|20)\d{2}\b|\b(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)"
        r"[a-z]*\.?\s+(?:19|20)\d{2}\b|\bpresent\b|\bcurrent\b",
        flags=re.IGNORECASE,
    )
    for line, start, end in lines:
        if date_pattern.search(line):
            evidence["employment_history"].append(_excerpt(text, start, end))
        lowered = line.casefold()
        if any(word in lowered for word in ("bachelor", "master", "ph.d", "phd", "degree", "diploma", "certification")):
            evidence["education_and_qualifications"].append(_excerpt(text, start, end))

    in_project_section = False
    for line, start, end in lines:
        lowered = line.strip().casefold().rstrip(":")
        if lowered in {"projects", "project experience", "selected projects"}:
            in_project_section = True
            continue
        if in_project_section and len(line.strip()) < 80 and lowered.endswith(":"):
            in_project_section = False
        if in_project_section:
            evidence["project_experience"].append(_excerpt(text, start, end))

    return ParsedCandidateCV(
        parser_version=PARSER_VERSION,
        employment_history=evidence["employment_history"],
        skills_and_technologies=evidence["skills_and_technologies"],
        education_and_qualifications=evidence["education_and_qualifications"],
        project_experience=evidence["project_experience"],
    )
